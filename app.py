from io import BytesIO
from pathlib import Path
import re
import time
import hmac

import pandas as pd
import streamlit as st
from settings import get_setting
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from kakao_route_service import (
    VEHICLE_TYPE_MAP,
    lookup_route,
    lookup_routes_from_df,
)
from result_history import (
    RESULTS_DIR,
    delete_result_history,
    get_result_history,
    initialize_result_history,
    remove_missing_file_records,
    save_result_file,
)

st.set_page_config(page_title="카카오 노선 조회기", layout="wide")
app_password = get_setting("APP_PASSWORD")
if not app_password:
    st.info("앱 접속 준비 중입니다. 관리자가 배포 설정에 APP_PASSWORD를 등록해야 합니다.")
    st.stop()
if not st.session_state.get("authenticated", False):
    st.title("카카오 노선 조회기")
    with st.form("login"):
        password = st.text_input("접속 비밀번호", type="password")
        submitted = st.form_submit_button("접속")
    if submitted:
        if hmac.compare_digest(password.encode(), app_password.encode()):
            st.session_state.authenticated = True
            st.rerun()
        else:
            st.error("비밀번호를 확인해 주세요.")
    st.stop()
if st.sidebar.button("로그아웃"):
    st.session_state.clear()
    st.rerun()
initialize_result_history()


def dataframe_to_excel_bytes(df: pd.DataFrame, sheet_name: str = "조회결과") -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
    return output.getvalue()


def write_dataframe_to_sheet(workbook, dataframe: pd.DataFrame, sheet_name: str) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    for row in dataframe_to_rows(dataframe, index=False, header=True):
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def original_workbook_with_results(
    original_file_bytes: bytes,
    result_df: pd.DataFrame,
    failure_df: pd.DataFrame,
) -> bytes:
    workbook = load_workbook(BytesIO(original_file_bytes))
    write_dataframe_to_sheet(workbook, result_df, "조회결과")
    write_dataframe_to_sheet(workbook, failure_df, "주소정제필요")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


SIDO_ALIAS = {
    "서울특별시": ["서울", "서울특별시", "서울시"],
    "부산광역시": ["부산", "부산광역시", "부산시"],
    "대구광역시": ["대구", "대구광역시", "대구시"],
    "인천광역시": ["인천", "인천광역시", "인천시"],
    "광주광역시": ["광주", "광주광역시", "광주시"],
    "대전광역시": ["대전", "대전광역시", "대전시"],
    "울산광역시": ["울산", "울산광역시", "울산시"],
    "세종특별자치시": ["세종", "세종시", "세종특별자치시"],
    "경기도": ["경기", "경기도"],
    "강원특별자치도": ["강원", "강원도", "강원특별자치도"],
    "충청북도": ["충북", "충청북도"],
    "충청남도": ["충남", "충청남도"],
    "전북특별자치도": ["전북", "전라북도", "전북특별자치도"],
    "전라남도": ["전남", "전라남도"],
    "경상북도": ["경북", "경상북도"],
    "경상남도": ["경남", "경상남도"],
    "제주특별자치도": ["제주", "제주도", "제주특별자치도"],
}
ALIAS_TO_STD = {
    alias: standard
    for standard, aliases in SIDO_ALIAS.items()
    for alias in aliases
}
SIDO_TO_GROUP = {
    "경기도": "수도권", "인천광역시": "수도권", "서울특별시": "수도권",
    "충청북도": "충청", "충청남도": "충청", "세종특별자치시": "충청", "대전광역시": "충청",
    "부산광역시": "경상", "대구광역시": "경상", "울산광역시": "경상", "경상남도": "경상", "경상북도": "경상",
    "광주광역시": "호남", "전북특별자치도": "호남", "전라남도": "호남",
    "강원특별자치도": "강원", "제주특별자치도": "제주",
}


def get_region(address):
    if pd.isna(address):
        return "", "", ""
    normalized_address = re.sub(r"\s+", " ", str(address).strip())
    tokens = normalized_address.split(" ")
    if not tokens:
        return "", "", ""
    raw_sido = tokens[0]
    sido = ALIAS_TO_STD.get(raw_sido, raw_sido)
    sigungu = ""
    if len(tokens) > 1 and tokens[1].endswith(("시", "군", "구")):
        sigungu = tokens[1]
    return sido, sigungu, SIDO_TO_GROUP.get(sido, "")


def add_region_columns(df: pd.DataFrame) -> pd.DataFrame:
    result_df = df.copy()
    origin_col = "출발지_조회주소" if "출발지_조회주소" in result_df.columns else "출발지"
    destination_col = "도착지_조회주소" if "도착지_조회주소" in result_df.columns else "도착지"
    origin_result = result_df[origin_col].apply(get_region)
    result_df["출발지_시도"] = origin_result.str[0]
    result_df["출발지_시군구"] = origin_result.str[1]
    result_df["출발지_대권역"] = origin_result.str[2]
    destination_result = result_df[destination_col].apply(get_region)
    result_df["도착지_시도"] = destination_result.str[0]
    result_df["도착지_시군구"] = destination_result.str[1]
    result_df["도착지_대권역"] = destination_result.str[2]
    return result_df


def build_failure_dataframe(result_df: pd.DataFrame) -> pd.DataFrame:
    failure_rows = []
    common_columns = [
        column for column in ["노선ID", "출발지코드", "도착지코드", "차종", "API상태"]
        if column in result_df.columns
    ]
    for _, row in result_df.iterrows():
        for endpoint in ["출발지", "도착지"]:
            status_column = f"{endpoint}_좌표상태"
            if row.get(status_column) != "실패":
                continue
            failure_row = {column: row.get(column) for column in common_columns}
            failure_row.update({
                "구분": endpoint,
                "원본주소": row.get(f"{endpoint}_원본주소", ""),
                "마지막_API검색어": row.get(f"{endpoint}_API검색어", ""),
                "정제단계": row.get(f"{endpoint}_정제단계", "조회실패"),
                "좌표상태": row.get(status_column, "실패"),
                "실패원인": row.get(f"{endpoint}_좌표메모", "주소검색 실패"),
                "조회소요시간_초": row.get(f"{endpoint}_조회소요시간_초", None),
            })
            failure_rows.append(failure_row)
    columns = common_columns + [
        "구분", "원본주소", "마지막_API검색어", "정제단계",
        "좌표상태", "실패원인", "조회소요시간_초",
    ]
    failure_df = pd.DataFrame(failure_rows, columns=columns)
    if not failure_df.empty:
        dedupe_columns = [column for column in ["구분", "원본주소", "차종"] if column in failure_df.columns]
        failure_df = failure_df.drop_duplicates(subset=dedupe_columns, keep="first").reset_index(drop=True)
    return failure_df


def show_history_page() -> None:
    st.title("조회 이력")
    st.caption("저장된 결과를 다시 내려받을 수 있습니다. 영구 저장소가 연결되지 않은 배포에서는 재배포 시 이력이 사라질 수 있으므로 결과를 다운로드해 보관하세요.")

    controls1, controls2 = st.columns([1, 4])
    with controls1:
        if st.button("목록 새로고침"):
            st.rerun()
    with controls2:
        if st.button("없는 파일 이력 정리"):
            removed = remove_missing_file_records()
            st.success(f"파일이 없는 이력 {removed}건을 정리했습니다.")
            st.rerun()

    history = get_result_history()
    if not history:
        st.info("저장된 조회 결과가 없습니다.")
        return

    for item in history:
        file_path = Path(item["result_file_path"])
        with st.container(border=True):
            left, middle, right = st.columns([5, 2, 2])
            with left:
                st.markdown(f"**{item['job_name']}**")
                st.caption(
                    f"{item['created_at']} · 원본: {item['original_file_name']} · "
                    f"결과: {item['result_file_name']}"
                )
            with middle:
                st.metric("조회 건수", f"{item['row_count']:,}")
                st.caption(f"성공 {item['success_count']:,} / 확인 필요 {item['failure_count']:,}")
            with right:
                if file_path.exists():
                    st.download_button(
                        "결과 다운로드",
                        data=file_path.read_bytes(),
                        file_name=item["result_file_name"],
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"download_{item['id']}",
                        width="stretch",
                    )
                else:
                    st.error("파일 없음")
                if st.button(
                    "이력·파일 삭제",
                    key=f"delete_{item['id']}",
                    width="stretch",
                ):
                    if delete_result_history(item["id"], delete_file=True):
                        st.success("삭제했습니다.")
                        st.rerun()
                    else:
                        st.error("삭제하지 못했습니다. 파일이 열려 있는지 확인하세요.")


def show_lookup_page() -> None:
    st.title("노선 기초 정보 조회 자동화 Test")
    st.caption("주소를 입력하거나 엑셀을 업로드하면 거리, 예상시간, 통행료를 조회합니다.")

    vehicle_options = list(VEHICLE_TYPE_MAP.keys())
    tab_manual, tab_excel = st.tabs(["수기 조회", "엑셀 업로드"])

    with tab_manual:
        st.subheader("수기 조회")
        column1, column2, column3 = st.columns([2, 2, 1])
        with column1:
            origin = st.text_input("출발지", placeholder="예: 경기도 의왕시 ...")
        with column2:
            destination = st.text_input("도착지", placeholder="예: 부산광역시 강서구 ...")
        with column3:
            vehicle = st.selectbox("조회할 차종", vehicle_options, index=vehicle_options.index("5톤"))

        if st.button("조회하기", type="primary"):
            if not origin or not destination:
                st.warning("출발지와 도착지를 모두 입력하세요.")
            else:
                with st.spinner("카카오 API 조회 중..."):
                    try:
                        result = lookup_route(origin, destination, vehicle)
                        result_df = pd.DataFrame([result])
                        st.dataframe(result_df, width="stretch")
                        if result.get("API상태") == "OK":
                            metric1, metric2, metric3 = st.columns(3)
                            metric1.metric("주행거리", f"{result['주행거리_km']} km")
                            metric2.metric("예상시간", f"{result['예상시간_분']} 분")
                            metric3.metric("톨비", f"{result['톨비_원']:,} 원")
                    except Exception as exc:
                        st.error(f"조회 중 오류가 발생했습니다: {type(exc).__name__} - {exc}")

    with tab_excel:
        st.subheader("엑셀 업로드 일괄 조회")
        st.info("엑셀에는 최소 '출발지', '도착지' 컬럼이 필요합니다. 화면에서 선택한 차종이 일괄 적용됩니다.")
        sample_df = pd.DataFrame({
            "노선ID": ["R001", "R002"],
            "출발지코드": ["O001", "O002"],
            "출발지": ["경기도 의왕시", "경기도 안성시"],
            "도착지코드": ["D001", "D002"],
            "도착지": ["부산광역시 강서구", "대구광역시 달성군"],
        })
        st.download_button(
            "샘플 엑셀 다운로드",
            data=dataframe_to_excel_bytes(sample_df, "주소목록"),
            file_name="sample_routes.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        uploaded_file = st.file_uploader("주소 엑셀 업로드", type=["xlsx"])
        if uploaded_file is None:
            return

        try:
            uploaded_bytes = uploaded_file.getvalue()
            excel_file = pd.ExcelFile(BytesIO(uploaded_bytes))
            sheet_name = st.selectbox("읽을 시트를 선택하세요", excel_file.sheet_names)
            input_df = pd.read_excel(BytesIO(uploaded_bytes), sheet_name=sheet_name)
            input_df.columns = input_df.columns.map(str).str.strip()

            job_name = st.text_input(
                "작업명",
                value=Path(uploaded_file.name).stem,
                help="조회 이력에 표시될 이름입니다.",
            )
            selected_vehicles = st.multiselect(
                "조회할 차종을 선택하세요",
                vehicle_options,
                default=["5톤"],
                help="여러 차종을 선택하면 각 노선을 차종별로 각각 조회합니다.",
            )
            st.write("업로드 미리보기")
            st.dataframe(input_df.head(20), width="stretch")

            if st.button("일괄 조회 시작", type="primary"):
                if not selected_vehicles:
                    st.warning("조회할 차종을 한 개 이상 선택하세요.")
                    return

                progress_bar = st.progress(0, text="조회 준비 중...")
                status_box = st.empty()
                detail_box = st.empty()
                started_at = time.perf_counter()

                def update_progress(completed: int, total: int, current: dict) -> None:
                    elapsed = time.perf_counter() - started_at
                    average = elapsed / completed if completed else 0
                    remaining = max(total - completed, 0) * average
                    percent = completed / total if total else 1.0
                    remaining_text = (
                        f"약 {int(remaining // 60)}분 {int(remaining % 60)}초"
                        if remaining >= 60 else f"약 {int(remaining)}초"
                    )
                    progress_bar.progress(percent, text=f"노선 조회 중 · {completed}/{total}건")
                    status_box.markdown(
                        f"**진행률:** {percent * 100:.1f}% | **예상 남은 시간:** {remaining_text}"
                    )
                    detail_box.caption(
                        f"현재 조회: {current.get('출발지', '')} → {current.get('도착지', '')} "
                        f"[{current.get('차종', '')}] · 상태: {current.get('API상태', '')}"
                    )

                result_df = lookup_routes_from_df(
                    input_df,
                    selected_vehicles=selected_vehicles,
                    progress_callback=update_progress,
                )
                result_df = add_region_columns(result_df)
                failure_df = build_failure_dataframe(result_df)
                result_file = original_workbook_with_results(uploaded_bytes, result_df, failure_df)

                success_count = int((result_df["API상태"] == "OK").sum())
                failure_count = len(failure_df)
                saved = save_result_file(
                    file_bytes=result_file,
                    original_file_name=uploaded_file.name,
                    row_count=len(result_df),
                    success_count=success_count,
                    failure_count=failure_count,
                    job_name=job_name,
                )

                total_elapsed = time.perf_counter() - started_at
                progress_bar.progress(1.0, text="조회 완료")
                status_box.markdown(
                    f"**완료:** {len(result_df)}건 · 총 소요시간 "
                    f"{int(total_elapsed // 60)}분 {int(total_elapsed % 60)}초"
                )
                detail_box.empty()
                st.success(
                    f"조회 완료: 정상 경로 {success_count}건, 주소 확인 필요 {failure_count}건\n\n"
                    f"결과 파일이 results 폴더에 자동 저장되었습니다."
                )
                st.code(saved["result_file_path"], language=None)
                st.dataframe(result_df, width="stretch")
                if not failure_df.empty:
                    st.warning("좌표를 찾지 못한 주소는 '주소정제필요' 시트에 정리했습니다.")
                    st.dataframe(failure_df, width="stretch")
                st.download_button(
                    "결과 엑셀 다운로드",
                    data=result_file,
                    file_name=saved["result_file_name"],
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        except Exception as exc:
            st.error(f"엑셀 처리 중 오류가 발생했습니다: {type(exc).__name__} - {exc}")


page = st.sidebar.radio(
    "메뉴",
    ["노선 조회", "조회 이력"],
    index=0,
)
st.sidebar.caption("카카오 노선 조회")

if page == "조회 이력":
    show_history_page()
else:
    show_lookup_page()
